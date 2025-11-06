#!/usr/bin/env python3
import subprocess
import yaml
import logging
import shutil
import sys
from pathlib import Path
from docxtpl import DocxTemplate
from jinja2 import Environment, StrictUndefined
from docx import Document
import time
from docxcompose.composer import Composer


class PDFGenerator:
    def __init__(self):
        logging.basicConfig(
            level=logging.INFO,
            format="%(asctime)s %(name)s %(levelname)s: %(message)s"
        )
        self.logger = logging.getLogger(self.__class__.__name__)

        self.repo_root = Path.cwd()
        self.mkdocs_yml = self.repo_root / "mkdocs.yml"
        self.doc_folder = self.repo_root / "docs"
        self.output_folder = self.repo_root / "public"
        self.tmp_folder = self.output_folder / "tmp"
        self.output_folder.mkdir(exist_ok=True)
        self.tmp_folder.mkdir(parents=True, exist_ok=True)

    def load_config(self):
        try:
            cfg = yaml.safe_load(self.mkdocs_yml.read_text())
        except Exception as e:
            self.logger.error("Failed to parse mkdocs.yml: %s", e)
            sys.exit(1)
        return cfg.get("nav", []), cfg.get("extra", {}).get("documents", {})

    def flatten_nav(self, nav_tree, depth=0):
        result = []
        for item in nav_tree:
            if isinstance(item, dict):
                for _, v in item.items():
                    if isinstance(v, str):
                        result.append((v, depth))
                    else:
                        result.extend(self.flatten_nav(v, depth + 1))
        return result

    def generate(self):
        nav, documents = self.load_config()
        jenv = Environment(undefined=StrictUndefined)

        for doc_id, meta in documents.items():
            self.logger.info("Processing '%s'", doc_id)
            nav_key = meta.get("nav_part") or self._error("Missing nav_part", doc_id)
            subtree = next((item[nav_key] for item in nav
                            if isinstance(item, dict) and nav_key in item),
                           None) or self._error(f"nav_part '{nav_key}' not found", doc_id)
            chapters = self.flatten_nav(subtree) or self._error(f"No chapters under '{nav_key}'", doc_id)

            doc_parts = []
            for i, (rel, depth) in enumerate(chapters):
                src = self.doc_folder / rel
                if not src.exists():
                    self._error(f"Missing source file: {src}", doc_id)
                dst = self.tmp_folder / f"{doc_id}_chap{i:02}.docx"
                self._pandoc(src, dst, max(depth - 1, 0), doc_id)
                doc_parts.append(dst)

            tpl_rel = meta.get("template") or self._error("Missing template path", doc_id)
            tpl_path = self.repo_root / tpl_rel
            if not tpl_path.exists():
                self._error(f"Template not found: {tpl_path}", doc_id)
            self.logger.info("[%s] Using template: %s", doc_id, tpl_path)

            tpl = DocxTemplate(str(tpl_path))
            placeholders = tpl.get_undeclared_template_variables()
            cfg_keys = set(meta.keys())
            self._check_keys(placeholders, cfg_keys, doc_id)

            try:
                tpl.render(meta)
            except Exception as e:
                self._error(f"Templating error: {e}", doc_id)

            filled = self.tmp_folder / f"{doc_id}_filled.docx"
            tpl.save(filled)

            master_doc = Document(str(filled))
            composer = Composer(master_doc)
            for part in doc_parts:
                composer.append(Document(str(part)))
            final_docx = self.tmp_folder / f"{doc_id}_combined.docx"
            composer.save(final_docx)

            self._export_pdf(final_docx, doc_id)

        shutil.rmtree(self.tmp_folder)
        self.logger.info("✅ All done.")

    def _check_keys(self, placeholders, cfg_keys, doc_id):
        missing = placeholders - cfg_keys
        unused = cfg_keys - placeholders
        if missing:
            self._error(f"Missing config keys: {missing}", doc_id)
        if unused:
            self.logger.warning("[%s] Config unused keys: %s", doc_id, unused)

    def _error(self, msg, doc_id=None):
        if doc_id:
            self.logger.error(f"[{doc_id}] {msg}")
        else:
            self.logger.error(msg)
        sys.exit(1)

    def _pandoc(self, src, dst, shift, doc_id):
        self.logger.info("[%s] Converting: %s → %s (shift=%d)", doc_id, src, dst, shift)
        try:
            subprocess.run([
                "pandoc", str(src),
                "--shift-heading-level-by", str(shift),
                "-f", "markdown", "-t", "docx", "-o", str(dst)
            ], check=True)
        except subprocess.CalledProcessError:
            self._error(f"Pandoc failed for {src}", doc_id)

    import time

    def _export_pdf(self, docx_file: Path, doc_id: str):
        try:
            if not docx_file.exists():
                self._error(f"File does not exist before export: {docx_file}", doc_id)
            if docx_file.stat().st_size == 0:
                self._error(f"File is empty before export: {docx_file}", doc_id)

            self.logger.info("[%s] Final DOCX size: %d bytes", doc_id, docx_file.stat().st_size)

            time.sleep(0.2)  # Ensure file is flushed

            # final PDF path under public/
            target_pdf = self.output_folder / f"{doc_id}_combined.pdf"

            subprocess.run([
                "libreoffice", "--headless", "--nologo",
                "--convert-to", "pdf", str(docx_file),
                "--outdir", str(self.output_folder)
            ], check=True)

            # make sure the generated file is where we expect
            if target_pdf.exists():
                self.logger.info("[%s] PDF saved: %s", doc_id, target_pdf)
            else:
                self._error(f"PDF not found at expected location: {target_pdf}", doc_id)

        except subprocess.CalledProcessError as e:
            self._error(f"PDF export failed: {e}", doc_id)
